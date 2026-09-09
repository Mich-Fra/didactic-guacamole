from flask import Flask, render_template, request, redirect, url_for, send_file
import datetime
import sqlite3
import holidays
import openpyxl
import calendar
import math

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO


app = Flask(__name__)

DB_PATH = "database.db"


# ============================================================
# COSTANTI
# ============================================================

GIORNI_SETTIMANA = [
    "Lunedì",
    "Martedì",
    "Mercoledì",
    "Giovedì",
    "Venerdì",
    "Sabato",
    "Domenica"
]

GIORNI_LAVORATIVI = [
    "Lunedì",
    "Martedì",
    "Mercoledì",
    "Giovedì",
    "Venerdì"
]

NOMI_MESI = [
    "Gennaio",
    "Febbraio",
    "Marzo",
    "Aprile",
    "Maggio",
    "Giugno",
    "Luglio",
    "Agosto",
    "Settembre",
    "Ottobre",
    "Novembre",
    "Dicembre"
]

# Lunedì 03/01/2000.
# Punto di riferimento assoluto della rotazione.
EPOCA_ROTazione = datetime.date(2000, 1, 3)


# ============================================================
# DATABASE
# ============================================================

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS ruoli (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT UNIQUE NOT NULL
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS dipendenti (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT UNIQUE NOT NULL,
            giorno_forzato TEXT DEFAULT 'Nessuno',
            ruolo TEXT DEFAULT 'Nessuno',
            rotazione_posizione INTEGER DEFAULT 0
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS impostazioni (
            chiave TEXT PRIMARY KEY,
            valore TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS festivita (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data TEXT UNIQUE NOT NULL,
            descrizione TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS tetto_periodi (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            data_inizio TEXT NOT NULL,
            data_fine TEXT NOT NULL,
            tetto_max INTEGER NOT NULL
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS eccezioni_dipendenti (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            id_dipendente INTEGER NOT NULL,
            data_inizio TEXT NOT NULL,
            data_fine TEXT NOT NULL,
            FOREIGN KEY (id_dipendente)
                REFERENCES dipendenti(id)
                ON DELETE CASCADE
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS limiti_ruoli (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ruolo TEXT UNIQUE NOT NULL,
            max_per_giorno INTEGER NOT NULL DEFAULT 1
        )
    """)

    # --------------------------------------------------------
    # MIGRAZIONE DIPENDENTI
    # --------------------------------------------------------

    cursor.execute("PRAGMA table_info(dipendenti)")
    colonne = [
        row[1]
        for row in cursor.fetchall()
    ]

    if "giorno_forzato" not in colonne:
        cursor.execute("""
            ALTER TABLE dipendenti
            ADD COLUMN giorno_forzato TEXT DEFAULT 'Nessuno'
        """)

    if "ruolo" not in colonne:
        cursor.execute("""
            ALTER TABLE dipendenti
            ADD COLUMN ruolo TEXT DEFAULT 'Nessuno'
        """)

    if "rotazione_posizione" not in colonne:
        cursor.execute("""
            ALTER TABLE dipendenti
            ADD COLUMN rotazione_posizione INTEGER DEFAULT 0
        """)

    # --------------------------------------------------------
    # RUOLI DEFAULT
    # --------------------------------------------------------

    cursor.execute(
        "SELECT COUNT(*) FROM ruoli"
    )

    if cursor.fetchone()[0] == 0:

        cursor.executemany(
            "INSERT INTO ruoli (nome) VALUES (?)",
            [
                ("Impiegato",),
                ("Team Leader",),
                ("Responsabile",),
                ("Manager",)
            ]
        )

    # --------------------------------------------------------
    # DIPENDENTI DEFAULT
    # --------------------------------------------------------

    cursor.execute(
        "SELECT COUNT(*) FROM dipendenti"
    )

    if cursor.fetchone()[0] == 0:

        cursor.executemany(
            """
            INSERT INTO dipendenti (
                nome,
                giorno_forzato,
                ruolo,
                rotazione_posizione
            )
            VALUES (?, ?, ?, ?)
            """,
            [
                ("Luigi Rossi", "Nessuno", "Nessuno", 0),
                ("Anna Bianchi", "Nessuno", "Nessuno", 1),
                ("Marco Neri", "Nessuno", "Nessuno", 2),
                ("Sofia Verdi", "Nessuno", "Nessuno", 3),
                ("Giovanni Gialli", "Nessuno", "Nessuno", 4)
            ]
        )

    # --------------------------------------------------------
    # FREQUENZA DEFAULT
    # --------------------------------------------------------

    cursor.execute("""
        SELECT COUNT(*)
        FROM impostazioni
        WHERE chiave = 'frequenza_rotazione'
    """)

    if cursor.fetchone()[0] == 0:

        cursor.execute("""
            INSERT INTO impostazioni (
                chiave,
                valore
            )
            VALUES (
                'frequenza_rotazione',
                '1'
            )
        """)

    conn.commit()
    conn.close()


init_db()


# ============================================================
# DATE
# ============================================================

def parse_date(value):

    if not value:
        return None

    value = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%d/%m/%Y"
    ):
        try:
            return datetime.datetime.strptime(
                value,
                fmt
            ).date()
        except ValueError:
            continue

    return None


def format_date_it(value):

    if isinstance(value, datetime.date):
        return value.strftime("%d/%m/%Y")

    return ""


# ============================================================
# NOMI
# ============================================================

def split_nome_cognome(nome_completo):

    parti = (
        nome_completo or ""
    ).strip().split(" ", 1)

    nome = (
        parti[0]
        if len(parti) > 0
        else ""
    )

    cognome = (
        parti[1]
        if len(parti) > 1
        else ""
    )

    return nome, cognome


# ============================================================
# NORMALIZZAZIONE
# ============================================================

def normalizza_frequenza(value):

    try:
        value = int(
            round(
                float(value)
            )
        )
    except (
        TypeError,
        ValueError
    ):
        value = 1

    return max(
        1,
        value
    )


# ============================================================
# CALCOLO SETTIMANA / CICLO
# ============================================================

def calcola_settimana_assoluta(
    lunedi
):
    return (
        lunedi - EPOCA_ROTazione
    ).days // 7


def calcola_ciclo_rotazione(
    lunedi,
    frequenza
):

    frequenza = normalizza_frequenza(
        frequenza
    )

    settimana_assoluta = (
        calcola_settimana_assoluta(
            lunedi
        )
    )

    return (
        settimana_assoluta
        // frequenza
    )


def calcola_inizio_ciclo(
    lunedi,
    frequenza
):

    frequenza = normalizza_frequenza(
        frequenza
    )

    ciclo = (
        calcola_ciclo_rotazione(
            lunedi,
            frequenza
        )
    )

    return (
        EPOCA_ROTazione
        + datetime.timedelta(
            weeks=ciclo * frequenza
        )
    )


def calcola_fine_ciclo(
    lunedi,
    frequenza
):

    inizio = calcola_inizio_ciclo(
        lunedi,
        frequenza
    )

    frequenza = normalizza_frequenza(
        frequenza
    )

    return (
        inizio
        + datetime.timedelta(
            weeks=frequenza,
            days=-1
        )
    )


# ============================================================
# FORMULE BIDIREZIONALI
# ============================================================

def persone_giorno_da_settimane(
    n_dipendenti,
    settimane
):
    """
    Numero minimo di persone necessarie al giorno
    per far passare tutti gli automatici entro X settimane.
    """

    try:
        settimane = int(
            settimane
        )
    except (
        TypeError,
        ValueError
    ):
        settimane = 1

    settimane = max(
        1,
        settimane
    )

    if n_dipendenti <= 0:
        return 1

    capacita_ciclo = (
        5 * settimane
    )

    return max(
        1,
        math.ceil(
            n_dipendenti
            / capacita_ciclo
        )
    )


def settimane_da_persone_giorno(
    n_dipendenti,
    persone_giorno
):
    """
    Numero minimo di settimane necessarie
    se abbiamo X persone disponibili al giorno.
    """

    try:
        persone_giorno = int(
            persone_giorno
        )
    except (
        TypeError,
        ValueError
    ):
        persone_giorno = 1

    persone_giorno = max(
        1,
        persone_giorno
    )

    if n_dipendenti <= 0:
        return 1

    capacita_settimanale = (
        5 * persone_giorno
    )

    return max(
        1,
        math.ceil(
            n_dipendenti
            / capacita_settimanale
        )
    )


def calcola_tetto_automatico(
    n_dipendenti,
    settimane
):

    return persone_giorno_da_settimane(
        n_dipendenti,
        settimane
    )


# ============================================================
# SLOT DI ROTAZIONE
# ============================================================

def calcola_settimana_target(
    posizione,
    ciclo,
    frequenza
):
    """
    Ogni ciclo contiene:

        5 giorni x frequenza settimane

    Ogni dipendente automatico ha un solo slot nel ciclo.

    Il giorno ruota di uno a ogni nuovo ciclo.
    """

    frequenza = normalizza_frequenza(
        frequenza
    )

    numero_slot = (
        5 * frequenza
    )

    posizione = int(
        posizione or 0
    )

    slot = (
        posizione % numero_slot
    )

    settimana_nel_ciclo, indice_giorno = (
        divmod(
            slot,
            5
        )
    )

    indice_giorno = (
        indice_giorno
        + ciclo
    ) % 5

    return (
        settimana_nel_ciclo,
        indice_giorno
    )


def calcola_data_target_rotazione(
    lunedi_settimana,
    frequenza,
    posizione
):

    frequenza = normalizza_frequenza(
        frequenza
    )

    ciclo = (
        calcola_ciclo_rotazione(
            lunedi_settimana,
            frequenza
        )
    )

    inizio_ciclo = (
        calcola_inizio_ciclo(
            lunedi_settimana,
            frequenza
        )
    )

    settimana_nel_ciclo, indice_giorno = (
        calcola_settimana_target(
            posizione,
            ciclo,
            frequenza
        )
    )

    return (
        inizio_ciclo
        + datetime.timedelta(
            weeks=settimana_nel_ciclo,
            days=indice_giorno
        )
    )


# ============================================================
# MIGRAZIONE FESTIVITÀ
# ============================================================

def normalizza_date_festivita():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id, data
        FROM festivita
    """)

    righe = cursor.fetchall()

    for (
        id_festa,
        data_raw
    ) in righe:

        dt = parse_date(
            data_raw
        )

        if dt is None:
            continue

        nuova_data = dt.strftime(
            "%Y-%m-%d"
        )

        if data_raw != nuova_data:

            try:

                cursor.execute("""
                    UPDATE festivita
                    SET data = ?
                    WHERE id = ?
                """, (
                    nuova_data,
                    id_festa
                ))

            except sqlite3.IntegrityError:

                cursor.execute("""
                    DELETE FROM festivita
                    WHERE id = ?
                """, (
                    id_festa,
                ))

    conn.commit()
    conn.close()


normalizza_date_festivita()


# ============================================================
# CARICA DATI
# ============================================================

def carica_dati():

    with get_db() as conn:

        cursor = conn.cursor()

        # ----------------------------------------------------
        # DIPENDENTI
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                id,
                nome,
                giorno_forzato,
                rotazione_posizione,
                ruolo
            FROM dipendenti
            ORDER BY
                rotazione_posizione ASC,
                id ASC
        """)

        dipendenti = cursor.fetchall()

        # ----------------------------------------------------
        # IMPOSTAZIONI
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                chiave,
                valore
            FROM impostazioni
        """)

        impostazioni = dict(
            cursor.fetchall()
        )

        # ----------------------------------------------------
        # FREQUENZA ROTAZIONE
        # ----------------------------------------------------

        frequenza = normalizza_frequenza(
            impostazioni.get(
                "frequenza_rotazione",
                1
            )
        )

        # ----------------------------------------------------
        # PROTEZIONE VENERDÌ → LUNEDÌ
        #
        # Se l'impostazione non esiste ancora,
        # il comportamento predefinito rimane ATTIVO.
        # Questo preserva il comportamento attuale.
        # ----------------------------------------------------

        valore_protezione = str(
            impostazioni.get(
                "protezione_venerdi_lunedi",
                "1"
            )
        ).strip().lower()

        protezione_venerdi_lunedi = (
            valore_protezione
            in (
                "1",
                "true",
                "yes",
                "on"
            )
        )

        # ----------------------------------------------------
        # AUTOMATICI
        # ----------------------------------------------------

        automatici = [
            d
            for d in dipendenti
            if (
                not d[2]
                or d[2] == "Nessuno"
            )
        ]

        automatici.sort(
            key=lambda x: (
                int(x[3] or 0),
                int(x[0])
            )
        )

        # ----------------------------------------------------
        # TETTO AUTOMATICO
        # ----------------------------------------------------

        tetto_auto = (
            calcola_tetto_automatico(
                len(automatici),
                frequenza
            )
        )

        # ----------------------------------------------------
        # TETTO MANUALE
        # ----------------------------------------------------

        try:

            tetto_manuale = int(
                impostazioni.get(
                    "tetto_manuale",
                    0
                ) or 0
            )

        except (
            TypeError,
            ValueError
        ):

            tetto_manuale = 0

        if tetto_manuale < 1:
            tetto_manuale = 0

        tetto_base = (
            tetto_manuale
            if tetto_manuale
            else tetto_auto
        )

        # ----------------------------------------------------
        # FESTIVITÀ PERSONALIZZATE
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                data,
                descrizione
            FROM festivita
        """)

        festivita = {}

        for (
            data_raw,
            descrizione
        ) in cursor.fetchall():

            dt = parse_date(
                data_raw
            )

            if dt:

                festivita[dt] = (
                    descrizione
                    or "Festività personalizzata"
                )

        # ----------------------------------------------------
        # ECCEZIONI
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                id_dipendente,
                data_inizio,
                data_fine
            FROM eccezioni_dipendenti
        """)

        eccezioni = {}

        for (
            id_dipendente,
            data_inizio,
            data_fine
        ) in cursor.fetchall():

            di = parse_date(
                data_inizio
            )

            df = parse_date(
                data_fine
            )

            if (
                di
                and df
            ):

                eccezioni.setdefault(
                    int(id_dipendente),
                    []
                ).append(
                    (
                        di,
                        df
                    )
                )

        # ----------------------------------------------------
        # PERIODI PERSONALIZZATI
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                id,
                data_inizio,
                data_fine,
                tetto_max
            FROM tetto_periodi
            ORDER BY
                data_inizio DESC,
                id DESC
        """)

        periodi = []

        for (
            id_periodo,
            data_inizio,
            data_fine,
            tetto
        ) in cursor.fetchall():

            di = parse_date(
                data_inizio
            )

            df = parse_date(
                data_fine
            )

            try:

                tetto = int(
                    tetto
                )

            except (
                TypeError,
                ValueError
            ):

                continue

            if (
                di is None
                or df is None
                or df < di
                or tetto < 1
            ):

                continue

            periodi.append({
                "id": int(
                    id_periodo
                ),
                "inizio": di,
                "fine": df,
                "tetto": tetto
            })

        # ----------------------------------------------------
        # LIMITI RUOLI
        # ----------------------------------------------------

        cursor.execute("""
            SELECT
                ruolo,
                max_per_giorno
            FROM limiti_ruoli
        """)

        limiti_ruoli = {}

        for (
            ruolo,
            massimo
        ) in cursor.fetchall():

            try:

                limiti_ruoli[
                    ruolo or "Nessuno"
                ] = max(
                    0,
                    int(massimo)
                )

            except (
                TypeError,
                ValueError
            ):

                continue

    return {
        "dipendenti": dipendenti,

        "automatici": automatici,

        "frequenza": frequenza,

        "tetto_auto": tetto_auto,

        "tetto_base": max(
            1,
            tetto_base
        ),

        "tetto_manuale": (
            tetto_manuale
            if tetto_manuale >= 1
            else None
        ),

        "festivita": festivita,

        "eccezioni": eccezioni,

        "periodi": periodi,

        "limiti_ruoli": limiti_ruoli,

        # ====================================================
        # NUOVA IMPOSTAZIONE
        # ====================================================

        "protezione_venerdi_lunedi":
            protezione_venerdi_lunedi
    }


# ============================================================
# PERIODO ATTIVO
# ============================================================

def tetto_per_data(
    data,
    tetto_base,
    periodi
):
    """
    Tetto effettivo della giornata.

    Se non c'è un periodo:
        tetto_base

    Se c'è un periodo:
        tetto del periodo

    Se esistono più periodi sovrapposti:
        prevale quello iniziato più recentemente.
    """

    validi = [
        periodo
        for periodo in periodi
        if (
            periodo["inizio"]
            <= data
            <= periodo["fine"]
        )
    ]

    if not validi:

        return max(
            1,
            int(tetto_base)
        )

    validi.sort(
        key=lambda p: (
            p["inizio"],
            p["id"]
        ),
        reverse=True
    )

    return max(
        1,
        int(
            validi[0]["tetto"]
        )
    )


def periodo_attivo(
    data,
    periodi
):

    return any(
        periodo["inizio"]
        <= data
        <= periodo["fine"]
        for periodo in periodi
    )


# ============================================================
# ECCEZIONI
# ============================================================

def in_eccezione(
    id_dipendente,
    data,
    eccezioni
):

    for (
        di,
        df
    ) in eccezioni.get(
        int(id_dipendente),
        []
    ):

        if (
            di
            <= data
            <= df
        ):
            return True

    return False


# ============================================================
# STATO GIORNALIERO
# ============================================================

def crea_stato_giorno(
    data,
    dati,
    festivita_nazionali
):

    is_weekend = (
        data.weekday() >= 5
    )

    is_festa_nazionale = (
        data in festivita_nazionali
    )

    is_festa_custom = (
        data in dati["festivita"]
    )

    if is_weekend:

        tipo = "Weekend"
        info = "Chiuso"

    elif (
        is_festa_nazionale
        or is_festa_custom
    ):

        tipo = "Festa"

        if is_festa_nazionale:

            info = (
                festivita_nazionali.get(
                    data,
                    "Festività nazionale"
                )
            )

        else:

            info = dati[
                "festivita"
            ].get(
                data,
                "Festività personalizzata"
            )

    else:

        tipo = "Lavoro"
        info = ""

    return {
        "data": data,
        "tipo": tipo,
        "info": info,

        # IMPORTANTE:
        # questo è il tetto realmente valido
        # per quella precisa data.
        "tetto": tetto_per_data(
            data,
            dati["tetto_base"],
            dati["periodi"]
        ),

        "periodo_forzato": periodo_attivo(
            data,
            dati["periodi"]
        ),

        "nomi": [],

        "dipendenti": [],

        "conteggio_ruoli": {}
    }


# ============================================================
# ASSEGNAZIONE
# ============================================================

def prova_assegnazione(
    stato,
    dipendente,
    dati
):

    id_dipendente = int(
        dipendente["id"]
    )

    nome = dipendente["nome"]

    ruolo = (
        dipendente["ruolo"]
        or "Nessuno"
    )

    # --------------------------------------------------------
    # NON LAVORATIVO
    # --------------------------------------------------------

    if stato["tipo"] != "Lavoro":
        return False

    # --------------------------------------------------------
    # DUPLICATO
    # --------------------------------------------------------

    if nome in stato["nomi"]:
        return False

    # --------------------------------------------------------
    # TETTO GIORNALIERO
    # --------------------------------------------------------

    if (
        len(stato["nomi"])
        >= stato["tetto"]
    ):
        return False

    # --------------------------------------------------------
    # LIMITE RUOLO
    # --------------------------------------------------------

    limite_ruolo = dati[
        "limiti_ruoli"
    ].get(
        ruolo
    )

    if limite_ruolo is not None:

        presenti_ruolo = (
            stato[
                "conteggio_ruoli"
            ].get(
                ruolo,
                0
            )
        )

        if (
            presenti_ruolo
            >= limite_ruolo
        ):
            return False

    # --------------------------------------------------------
    # INSERIMENTO
    # --------------------------------------------------------

    stato["nomi"].append(
        nome
    )

    stato["dipendenti"].append({
        "id": id_dipendente,
        "nome": nome,
        "ruolo": ruolo
    })

    stato[
        "conteggio_ruoli"
    ][ruolo] = (
        stato[
            "conteggio_ruoli"
        ].get(
            ruolo,
            0
        ) + 1
    )

    return True


# ============================================================
# PROTEZIONE VENERDÌ -> LUNEDÌ
# ============================================================

def applica_protezione_venerdi_lunedi(
    stato,
    automatici,
    dati,
    festivita_nazionali,
    inizio_ciclo,
    stato_precedente=None
):
    """
    Protegge il lunedì successivo al venerdì senza ricostruire
    la rotazione da zero.

    Regola:
        - un automatico che ha lavorato il venerdì precedente
          non deve rimanere assegnato al lunedì successivo,
          quando esiste un altro automatico con cui effettuare
          uno scambio valido;
        - il numero di persone sul lunedì resta invariato;
        - si modifica il minimo indispensabile: normalmente
          l'automatico del lunedì viene spostato al martedì e
          l'automatico del martedì prende il suo posto.

    NOTA IMPORTANTE:
    non è matematicamente possibile mantenere contemporaneamente,
    per tutti i dipendenti, un avanzamento esatto di un giorno
    ad ogni ciclo e vietare sempre Venerdì -> Lunedì. In quel caso
    il motore usa quindi lo scambio minimo necessario, preservando
    la copertura e la rotazione il più possibile.
    """

    # Tutti i venerdì che precedono i lunedì presenti nel ciclo.
    venerdi_precedenti = {}
    lunedi_correnti = []

    for giorno in sorted(stato):
        if giorno.weekday() == 0:  # lunedì
            lunedi_correnti.append(giorno)
            venerdi_precedenti[giorno] = giorno - datetime.timedelta(days=3)

    if not lunedi_correnti:
        return

    auto_by_id = {
        int(c["id"]): c for c in automatici
    }

    def ids_del_giorno(giorno):
        return {
            int(x["id"])
            for x in stato[giorno]["dipendenti"]
            if int(x["id"]) in auto_by_id
        }

    def ha_lavorato_venerdi_precedente(candidato, lunedi):
        venerdi = venerdi_precedenti[lunedi]
        fonte = stato if venerdi in stato else (stato_precedente or {})
        if venerdi in fonte:
            return int(candidato["id"]) in {
                int(x["id"])
                for x in fonte[venerdi]["dipendenti"]
            }
        return False

    def rimuovi(stato_giorno, candidato):
        cid = int(candidato["id"])
        stato_giorno["dipendenti"] = [
            x for x in stato_giorno["dipendenti"]
            if int(x["id"]) != cid
        ]
        stato_giorno["nomi"] = [
            x["nome"]
            for x in stato_giorno["dipendenti"]
        ]
        ruolo = candidato.get("ruolo") or "Nessuno"
        count = stato_giorno["conteggio_ruoli"].get(ruolo, 0)
        if count > 1:
            stato_giorno["conteggio_ruoli"][ruolo] = count - 1
        else:
            stato_giorno["conteggio_ruoli"].pop(ruolo, None)

    def aggiungi(stato_giorno, candidato):
        ruolo = candidato.get("ruolo") or "Nessuno"
        stato_giorno["dipendenti"].append({
            "id": int(candidato["id"]),
            "nome": candidato["nome"],
            "ruolo": ruolo
        })
        stato_giorno["nomi"] = [
            x["nome"] for x in stato_giorno["dipendenti"]
        ]
        stato_giorno["conteggio_ruoli"][ruolo] = (
            stato_giorno["conteggio_ruoli"].get(ruolo, 0) + 1
        )

    def ruolo_ok(giorno, candidato):
        ruolo = candidato.get("ruolo") or "Nessuno"
        limite = dati["limiti_ruoli"].get(ruolo)
        if limite is None:
            return True
        return (
            stato[giorno]["conteggio_ruoli"].get(ruolo, 0) < limite
        )

    for lunedi in lunedi_correnti:
        if stato[lunedi]["tipo"] != "Lavoro":
            continue

        # Gli automatici che hanno lavorato il venerdì precedente
        # e sono attualmente sul lunedì.
        offensori = [
            auto_by_id[cid]
            for cid in ids_del_giorno(lunedi)
            if ha_lavorato_venerdi_precedente(auto_by_id[cid], lunedi)
        ]

        for offensore in offensori:
            # Potrebbe essere già stato spostato da uno scambio precedente.
            if int(offensore["id"]) not in ids_del_giorno(lunedi):
                continue
            if not ha_lavorato_venerdi_precedente(offensore, lunedi):
                continue

            # Cerchiamo prima martedì, poi mercoledì/giovedì/venerdì.
            giorni_candidati = []
            for delta in (1, 2, 3, 4):
                g = lunedi + datetime.timedelta(days=delta)
                if g not in stato:
                    continue
                if stato[g]["tipo"] != "Lavoro":
                    continue
                giorni_candidati.append(g)

            partner_scelto = None
            giorno_partner = None
            punteggio_migliore = None

            for giorno_partner_cand in giorni_candidati:
                for elemento in list(stato[giorno_partner_cand]["dipendenti"]):
                    cid = int(elemento["id"])
                    partner = auto_by_id.get(cid)
                    if partner is None:
                        continue

                    # Non trasferiamo al lunedì chi ha lavorato
                    # il venerdì precedente: altrimenti creiamo
                    # esattamente il problema che stiamo evitando.
                    if ha_lavorato_venerdi_precedente(partner, lunedi):
                        continue

                    # Eccezioni personali.
                    if in_eccezione(
                        partner["id"],
                        lunedi,
                        dati["eccezioni"]
                    ):
                        continue

                    # Verifica ruoli dopo aver tolto entrambi i candidati.
                    ruolo_off = offensore.get("ruolo") or "Nessuno"
                    ruolo_partner = partner.get("ruolo") or "Nessuno"

                    limite_off = dati["limiti_ruoli"].get(ruolo_off)
                    limite_partner = dati["limiti_ruoli"].get(ruolo_partner)

                    conteggi_lun = dict(stato[lunedi]["conteggio_ruoli"])
                    conteggi_partner = dict(stato[giorno_partner_cand]["conteggio_ruoli"])

                    conteggi_lun[ruolo_off] = conteggi_lun.get(ruolo_off, 0) - 1
                    conteggi_lun[ruolo_partner] = conteggi_lun.get(ruolo_partner, 0) + 1
                    conteggi_partner[ruolo_partner] = conteggi_partner.get(ruolo_partner, 0) - 1
                    conteggi_partner[ruolo_off] = conteggi_partner.get(ruolo_off, 0) + 1

                    if limite_partner is not None and conteggi_lun.get(ruolo_partner, 0) > limite_partner:
                        continue
                    if limite_off is not None and conteggi_partner.get(ruolo_off, 0) > limite_off:
                        continue

                    # Preferenza assoluta al martedì e poi alla minima
                    # distanza dal target teorico.
                    distanza = abs(
                        (giorno_partner_cand - offensore["data_target"]).days
                    )
                    distanza_partner = abs(
                        (lunedi - partner["data_target"]).days
                    )
                    punteggio = (
                        giorno_partner_cand.weekday(),
                        distanza + distanza_partner,
                        distanza,
                        distanza_partner,
                        partner["posizione"],
                        partner["id"]
                    )

                    if punteggio_migliore is None or punteggio < punteggio_migliore:
                        punteggio_migliore = punteggio
                        partner_scelto = partner
                        giorno_partner = giorno_partner_cand

            if partner_scelto is None:
                # Nessuno scambio sicuro disponibile. Il lunedì rimane
                # comunque coperto; non forziamo una violazione peggiore.
                continue

            # Scambio atomico: stessa capacità giornaliera, nessun buco.
            rimuovi(stato[lunedi], offensore)
            rimuovi(stato[giorno_partner], partner_scelto)

            aggiungi(stato[lunedi], partner_scelto)
            aggiungi(stato[giorno_partner], offensore)

            offensore["data_assegnata"] = giorno_partner
            partner_scelto["data_assegnata"] = lunedi


# ============================================================
# GENERAZIONE CICLO
# ============================================================

def genera_ciclo(
    inizio_ciclo,
    fine_ciclo,
    dati,
    festivita_nazionali,
    stato_precedente=None
):
    """
    Genera un ciclo completo.

    PRIORITÀ:

    1. Dipendenti con giorno fisso
    2. Target automatici
    3. Fallback
    4. Riequilibrio verso i periodi forzati

    IMPORTANTE:

    "Forza per periodo" non cambia la rotazione teorica.
    Cambia la capacità disponibile e permette al motore
    di spostare i turni automatici verso quelle date.
    """

    frequenza = dati["frequenza"]

    # ========================================================
    # CREA TUTTI I GIORNI DEL CICLO
    # ========================================================

    stato = {}

    data = inizio_ciclo

    while data <= fine_ciclo:

        stato[data] = crea_stato_giorno(
            data,
            dati,
            festivita_nazionali
        )

        data += datetime.timedelta(
            days=1
        )

    # ========================================================
    # DIPENDENTI FISSI
    # ========================================================

    for dipendente in dati["dipendenti"]:

        giorno_forzato = (
            dipendente[2]
            or "Nessuno"
        )

        if giorno_forzato == "Nessuno":
            continue

        if (
            giorno_forzato
            not in GIORNI_LAVORATIVI
        ):
            continue

        indice_giorno = (
            GIORNI_LAVORATIVI.index(
                giorno_forzato
            )
        )

        data = inizio_ciclo

        while data <= fine_ciclo:

            if (
                data.weekday()
                == indice_giorno
            ):

                if not in_eccezione(
                    dipendente[0],
                    data,
                    dati["eccezioni"]
                ):

                    prova_assegnazione(
                        stato[data],
                        {
                            "id": dipendente[0],
                            "nome": dipendente[1],
                            "ruolo": dipendente[4]
                        },
                        dati
                    )

            data += datetime.timedelta(
                days=1
            )

    # ========================================================
    # PREPARAZIONE AUTOMATICI
    # ========================================================

    ciclo = calcola_ciclo_rotazione(
        inizio_ciclo,
        frequenza
    )

    automatici = []

    for dipendente in dati["automatici"]:

        id_dipendente = int(
            dipendente[0]
        )

        posizione = int(
            dipendente[3]
            or 0
        )

        automatici.append({
            "id": id_dipendente,
            "nome": dipendente[1],
            "ruolo": (
                dipendente[4]
                or "Nessuno"
            ),
            "posizione": posizione,
            "data_target": None,
            "data_assegnata": None
        })

        (
            settimana_nel_ciclo,
            indice_giorno
        ) = calcola_settimana_target(
            posizione,
            ciclo,
            frequenza
        )

        data_target = (
            inizio_ciclo
            + datetime.timedelta(
                weeks=settimana_nel_ciclo,
                days=indice_giorno
            )
        )

        automatici[-1][
            "data_target"
        ] = data_target

    # --------------------------------------------------------
    # Ordine stabile
    # --------------------------------------------------------

    automatici.sort(
        key=lambda x: (
            x["data_target"],
            x["posizione"],
            x["id"]
        )
    )

    non_assegnati = []

    # ========================================================
    # TARGET
    # ========================================================

    for candidato in automatici:

        data_target = (
            candidato["data_target"]
        )

        if data_target not in stato:

            non_assegnati.append(
                candidato
            )
            continue

        if in_eccezione(
            candidato["id"],
            data_target,
            dati["eccezioni"]
        ):

            non_assegnati.append(
                candidato
            )
            continue

        if prova_assegnazione(
            stato[data_target],
            candidato,
            dati
        ):

            candidato[
                "data_assegnata"
            ] = data_target

        else:

            non_assegnati.append(
                candidato
            )

    # ========================================================
    # FALLBACK
    # ========================================================

    giorni_lavorativi = [
        giorno
        for giorno in sorted(
            stato.keys()
        )
        if (
            stato[giorno]["tipo"]
            == "Lavoro"
        )
    ]

    for candidato in non_assegnati:

        disponibili = []

        for giorno in giorni_lavorativi:

            if in_eccezione(
                candidato["id"],
                giorno,
                dati["eccezioni"]
            ):
                continue

            if (
                len(
                    stato[giorno]["nomi"]
                )
                >= stato[giorno]["tetto"]
            ):
                continue

            disponibili.append(
                giorno
            )

        # ----------------------------------------------------
        # Priorità:
        #
        # 1. periodo forzato
        # 2. distanza dal target
        # 3. data
        # ----------------------------------------------------

        disponibili.sort(
            key=lambda giorno: (
                0
                if stato[giorno][
                    "periodo_forzato"
                ]
                else 1,

                abs(
                    (
                        giorno
                        - candidato["data_target"]
                    ).days
                ),

                giorno
            )
        )

        for giorno in disponibili:

            if prova_assegnazione(
                stato[giorno],
                candidato,
                dati
            ):

                candidato[
                    "data_assegnata"
                ] = giorno

                break

    # ========================================================
    # RIEQUILIBRIO FORZA PER PERIODO
    # ========================================================
    #
    # Questa è la parte che fa funzionare realmente
    # "Forza per periodo".
    #
    # Esempio:
    #
    # normale = 2
    # periodo = 4
    #
    # il sistema può spostare turni automatici già assegnati
    # verso le giornate del periodo, fino alla nuova capacità.
    #
    # Non crea un secondo turno per il dipendente.
    # Sposta il turno esistente.
    # ========================================================

    giorni_forzati = [
        giorno
        for giorno in giorni_lavorativi
        if stato[giorno][
            "periodo_forzato"
        ]
    ]

    # --------------------------------------------------------
    # Prima le giornate forzate con più spazio.
    # --------------------------------------------------------

    giorni_forzati.sort(
        key=lambda giorno: (
            -(
                stato[giorno]["tetto"]
                - len(
                    stato[giorno]["nomi"]
                )
            ),
            giorno
        )
    )

    for giorno_forzato in giorni_forzati:

        while True:

            capacita = (
                stato[
                    giorno_forzato
                ]["tetto"]
            )

            occupati = len(
                stato[
                    giorno_forzato
                ]["nomi"]
            )

            if occupati >= capacita:
                break

            candidati_spostamento = []

            # ------------------------------------------------
            # Cerchiamo automatici già assegnati altrove.
            # ------------------------------------------------

            for candidato in automatici:

                data_attuale = (
                    candidato["data_assegnata"]
                )

                if data_attuale is None:
                    continue

                if (
                    data_attuale
                    == giorno_forzato
                ):
                    continue

                # --------------------------------------------
                # Se il dipendente non può andare in
                # quel giorno, non è candidato.
                # --------------------------------------------

                if in_eccezione(
                    candidato["id"],
                    giorno_forzato,
                    dati["eccezioni"]
                ):
                    continue

                ruolo = (
                    candidato["ruolo"]
                    or "Nessuno"
                )

                limite_ruolo = (
                    dati[
                        "limiti_ruoli"
                    ].get(
                        ruolo
                    )
                )

                if limite_ruolo is not None:

                    presenti_ruolo = (
                        stato[
                            giorno_forzato
                        ]["conteggio_ruoli"].get(
                            ruolo,
                            0
                        )
                    )

                    if (
                        presenti_ruolo
                        >= limite_ruolo
                    ):
                        continue

                # --------------------------------------------
                # Vogliamo spostare per primi i turni
                # più vicini alla data forzata.
                # --------------------------------------------

                distanza = abs(
                    (
                        giorno_forzato
                        - candidato["data_target"]
                    ).days
                )

                distanza_dal_corrente = abs(
                    (
                        giorno_forzato
                        - data_attuale
                    ).days
                )

                candidati_spostamento.append(
                    (
                        distanza,
                        distanza_dal_corrente,
                        candidato["posizione"],
                        candidato["id"],
                        candidato
                    )
                )

            if not candidati_spostamento:
                break

            candidati_spostamento.sort(
                key=lambda x: (
                    x[0],
                    x[1],
                    x[2],
                    x[3]
                )
            )

            spostato = False

            for (
                _,
                _,
                _,
                _,
                candidato
            ) in candidati_spostamento:

                vecchia_data = (
                    candidato["data_assegnata"]
                )

                nome = candidato[
                    "nome"
                ]

                ruolo = (
                    candidato["ruolo"]
                    or "Nessuno"
                )

                if vecchia_data not in stato:
                    continue

                stato_vecchio = (
                    stato[vecchia_data]
                )

                # --------------------------------------------
                # Rimuovi dalla giornata precedente
                # --------------------------------------------

                if nome not in (
                    stato_vecchio["nomi"]
                ):
                    continue

                stato_vecchio[
                    "nomi"
                ].remove(
                    nome
                )

                stato_vecchio[
                    "dipendenti"
                ] = [
                    elemento
                    for elemento in (
                        stato_vecchio[
                            "dipendenti"
                        ]
                    )
                    if elemento["id"]
                    != candidato["id"]
                ]

                ruolo_count = (
                    stato_vecchio[
                        "conteggio_ruoli"
                    ].get(
                        ruolo,
                        0
                    )
                )

                if ruolo_count > 1:

                    stato_vecchio[
                        "conteggio_ruoli"
                    ][ruolo] = (
                        ruolo_count - 1
                    )

                else:

                    stato_vecchio[
                        "conteggio_ruoli"
                    ].pop(
                        ruolo,
                        None
                    )

                # --------------------------------------------
                # Prova ad aggiungere al periodo forzato
                # --------------------------------------------

                riuscito = prova_assegnazione(
                    stato[
                        giorno_forzato
                    ],
                    candidato,
                    dati
                )

                if riuscito:

                    candidato[
                        "data_assegnata"
                    ] = giorno_forzato

                    spostato = True

                    break

                # --------------------------------------------
                # Se fallisce ripristina
                # --------------------------------------------

                prova_assegnazione(
                    stato_vecchio,
                    candidato,
                    dati
                )

            if not spostato:
                break

    # ========================================================
    # PROTEZIONE VENERDÌ → LUNEDÌ
    # ========================================================

    if dati.get(
        "protezione_venerdi_lunedi",
        True
    ):

        applica_protezione_venerdi_lunedi(
            stato,
            automatici,
            dati,
            festivita_nazionali,
            inizio_ciclo,
            stato_precedente
        )

    # ========================================================
    # ORDINA I NOMI
    # ========================================================

    for giorno in stato:

        stato[
            giorno
        ]["dipendenti"].sort(
            key=lambda x: (
                x["nome"].lower()
            )
        )

        stato[
            giorno
        ]["nomi"] = [
            elemento["nome"]
            for elemento in (
                stato[
                    giorno
                ]["dipendenti"]
            )
        ]

    return stato


# ============================================================
# CALENDARIO COMPLETO
# ============================================================

def ottieni_calendario_settimanale(
    data_inizio,
    data_fine
):

    dati = carica_dati()

    frequenza = dati[
        "frequenza"
    ]

    # --------------------------------------------------------
    # Estensione al lunedì/domenica
    # --------------------------------------------------------

    lunedi_inizio = (
        data_inizio
        - datetime.timedelta(
            days=data_inizio.weekday()
        )
    )

    domenica_fine = (
        data_fine
        + datetime.timedelta(
            days=6 - data_fine.weekday()
        )
    )

    # --------------------------------------------------------
    # Festività italiane
    # --------------------------------------------------------

    anni = set(
        range(
            lunedi_inizio.year,
            domenica_fine.year + 1
        )
    )

    festivita_nazionali = holidays.Italy(
        years=anni,
        language="it"
    )

    # --------------------------------------------------------
    # Cicli da generare
    # --------------------------------------------------------

    primo_ciclo = calcola_ciclo_rotazione(
        lunedi_inizio,
        frequenza
    )

    ultimo_ciclo = calcola_ciclo_rotazione(
        domenica_fine,
        frequenza
    )

    cicli = {}
    stato_precedente = None

    for numero_ciclo in range(
        primo_ciclo,
        ultimo_ciclo + 1
    ):

        inizio_ciclo = (
            EPOCA_ROTazione
            + datetime.timedelta(
                weeks=(
                    numero_ciclo
                    * frequenza
                )
            )
        )

        fine_ciclo = (
            inizio_ciclo
            + datetime.timedelta(
                weeks=frequenza,
                days=-1
            )
        )

        cicli[
            numero_ciclo
        ] = genera_ciclo(
            inizio_ciclo,
            fine_ciclo,
            dati,
            festivita_nazionali,
            stato_precedente
        )

        stato_precedente = cicli[numero_ciclo]

    # --------------------------------------------------------
    # Statistiche
    # --------------------------------------------------------

    statistiche = {
        dipendente[1]: 0
        for dipendente in dati["dipendenti"]
    }

    calendario = {}

    data = lunedi_inizio

    while data <= domenica_fine:

        lunedi = (
            data
            - datetime.timedelta(
                days=data.weekday()
            )
        )

        iso = lunedi.isocalendar()

        nome_settimana = (
            f"Sett. {iso.week}"
        )

        calendario.setdefault(
            nome_settimana,
            []
        )

        # ----------------------------------------------------
        # WEEKEND
        # ----------------------------------------------------

        is_weekend = (
            data.weekday() >= 5
        )

        # ----------------------------------------------------
        # FESTA
        # ----------------------------------------------------

        is_festa_nazionale = (
            data in festivita_nazionali
        )

        is_festa_custom = (
            data in dati["festivita"]
        )

        giorno = {
            "data": data.strftime(
                "%d/%m/%Y"
            ),
            "giorno": GIORNI_SETTIMANA[
                data.weekday()
            ],
            "tipo": "Lavoro",
            "info": "",
            "lavoratori": []
        }

        if is_weekend:

            giorno.update({
                "tipo": "Weekend",
                "info": "Chiuso"
            })

        elif (
            is_festa_nazionale
            or is_festa_custom
        ):

            if is_festa_nazionale:

                nome_festa = (
                    festivita_nazionali.get(
                        data,
                        "Festività nazionale"
                    )
                )

            else:

                nome_festa = dati[
                    "festivita"
                ].get(
                    data,
                    "Festività personalizzata"
                )

            giorno.update({
                "tipo": "Festa",
                "info": nome_festa
            })

        else:

            numero_ciclo = (
                calcola_ciclo_rotazione(
                    lunedi,
                    frequenza
                )
            )

            stato_ciclo = (
                cicli.get(
                    numero_ciclo
                )
            )

            if stato_ciclo:

                stato_giorno = (
                    stato_ciclo.get(
                        data
                    )
                )

                if stato_giorno:

                    giorno[
                        "lavoratori"
                    ] = list(
                        stato_giorno[
                            "nomi"
                        ]
                    )

                    if (
                        data_inizio
                        <= data
                        <= data_fine
                    ):

                        for nome in (
                            stato_giorno["nomi"]
                        ):

                            if nome in statistiche:

                                statistiche[
                                    nome
                                ] += 1

        calendario[
            nome_settimana
        ].append(
            giorno
        )

        data += datetime.timedelta(
            days=1
        )

    return (
        calendario,
        statistiche,
        dati["tetto_base"]
    )


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    oggi = datetime.date.today()

    mese_param = request.args.get(
        "mese",
        "seleziona"
    )

    try:

        if mese_param != "seleziona":

            mese = int(
                mese_param
            )

            if not 1 <= mese <= 12:
                raise ValueError

            dt_start = datetime.date(
                oggi.year,
                mese,
                1
            )

            dt_end = datetime.date(
                oggi.year,
                mese,
                calendar.monthrange(
                    oggi.year,
                    mese
                )[1]
            )

        else:

            start = request.args.get(
                "start"
            )

            end = request.args.get(
                "end"
            )

            if start and end:

                dt_start = parse_date(
                    start
                )

                dt_end = parse_date(
                    end
                )

                if (
                    dt_start is None
                    or dt_end is None
                    or dt_end < dt_start
                ):
                    raise ValueError

            else:

                dt_start = datetime.date(
                    oggi.year,
                    oggi.month,
                    1
                )

                dt_end = datetime.date(
                    oggi.year,
                    oggi.month,
                    calendar.monthrange(
                        oggi.year,
                        oggi.month
                    )[1]
                )

    except (
        ValueError,
        TypeError
    ):

        dt_start = datetime.date(
            oggi.year,
            oggi.month,
            1
        )

        dt_end = datetime.date(
            oggi.year,
            oggi.month,
            calendar.monthrange(
                oggi.year,
                oggi.month
            )[1]
        )

    prev_mese = (
        12
        if dt_start.month == 1
        else dt_start.month - 1
    )

    prev_anno = (
        dt_start.year - 1
        if dt_start.month == 1
        else dt_start.year
    )

    next_mese = (
        1
        if dt_start.month == 12
        else dt_start.month + 1
    )

    next_anno = (
        dt_start.year + 1
        if dt_start.month == 12
        else dt_start.year
    )

    (
        calendario,
        statistiche,
        tetto_max
    ) = ottieni_calendario_settimanale(
        dt_start,
        dt_end
    )

    return render_template(
        "index.html",
        calendario=calendario,
        statistiche=statistiche,
        start_val=dt_start.strftime(
            "%Y-%m-%d"
        ),
        end_val=dt_end.strftime(
            "%Y-%m-%d"
        ),
        settimana_attiva=request.args.get(
            "settimana",
            "tutte"
        ),
        mesi_lista=NOMI_MESI,
        mese_attivo=mese_param,
        tetto_max=tetto_max,

        prev_start=datetime.date(
            prev_anno,
            prev_mese,
            1
        ).strftime(
            "%Y-%m-%d"
        ),

        prev_end=datetime.date(
            prev_anno,
            prev_mese,
            calendar.monthrange(
                prev_anno,
                prev_mese
            )[1]
        ).strftime(
            "%Y-%m-%d"
        ),

        next_start=datetime.date(
            next_anno,
            next_mese,
            1
        ).strftime(
            "%Y-%m-%d"
        ),

        next_end=datetime.date(
            next_anno,
            next_mese,
            calendar.monthrange(
                next_anno,
                next_mese
            )[1]
        ).strftime(
            "%Y-%m-%d"
        ),

        titolo_mese=(
            f"{NOMI_MESI[dt_start.month - 1]} "
            f"{dt_start.year}"
        )
    )


# ============================================================
# EXCEL
# ============================================================

@app.route("/scarica_excel")
def scarica_excel():

    start_param = request.args.get(
        "start"
    )

    end_param = request.args.get(
        "end"
    )

    settimana_attiva = request.args.get(
        "settimana",
        "tutte"
    )

    dt_start = parse_date(
        start_param
    )

    dt_end = parse_date(
        end_param
    )

    if (
        dt_start is None
        or dt_end is None
        or dt_end < dt_start
    ):

        oggi = datetime.date.today()

        dt_start = datetime.date(
            oggi.year,
            oggi.month,
            1
        )

        dt_end = datetime.date(
            oggi.year,
            oggi.month,
            calendar.monthrange(
                oggi.year,
                oggi.month
            )[1]
        )

        start_param = dt_start.strftime(
            "%Y-%m-%d"
        )

        end_param = dt_end.strftime(
            "%Y-%m-%d"
        )

    (
        calendario,
        _,
        _
    ) = ottieni_calendario_settimanale(
        dt_start,
        dt_end
    )

    wb = openpyxl.Workbook()

    ws = wb.active

    ws.title = "Piano Homeworking"

    font_titolo = Font(
        name="Segoe UI",
        size=11,
        bold=True,
        color="FFFFFF"
    )

    font_testo = Font(
        name="Segoe UI",
        size=10,
        color="000000"
    )

    font_festa = Font(
        name="Segoe UI",
        size=10,
        bold=True,
        color="9C0006"
    )

    font_weekend = Font(
        name="Segoe UI",
        size=10,
        color="7F7F7F"
    )

    fill_header = PatternFill(
        start_color="1A365D",
        end_color="1A365D",
        fill_type="solid"
    )

    fill_zebra = PatternFill(
        start_color="F8FAFC",
        end_color="F8FAFC",
        fill_type="solid"
    )

    fill_weekend = PatternFill(
        start_color="E2E8F0",
        end_color="E2E8F0",
        fill_type="solid"
    )

    fill_festa = PatternFill(
        start_color="FFC7CE",
        end_color="FFC7CE",
        fill_type="solid"
    )

    allineamento = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    bordo = Border(
        left=Side(
            style="thin",
            color="CBD5E1"
        ),
        right=Side(
            style="thin",
            color="CBD5E1"
        ),
        top=Side(
            style="thin",
            color="CBD5E1"
        ),
        bottom=Side(
            style="thin",
            color="CBD5E1"
        )
    )

    ws.append([
        "Settimana",
        "Data",
        "Giorno",
        "Stato / Persone in Homeworking"
    ])

    ws.row_dimensions[1].height = 28

    for cell in ws[1]:

        cell.font = font_titolo
        cell.fill = fill_header
        cell.alignment = allineamento
        cell.border = bordo

    riga_idx = 2

    for (
        nome_sett,
        giorni
    ) in calendario.items():

        if (
            settimana_attiva != "tutte"
            and settimana_attiva != nome_sett
        ):
            continue

        for giorno in giorni:

            if giorno["tipo"] == "Weekend":

                stato = (
                    "Chiuso (Weekend)"
                )

                fill = fill_weekend
                font = font_weekend

            elif giorno["tipo"] == "Festa":

                stato = (
                    f"🎉 Festa: "
                    f"{giorno['info']}"
                )

                fill = fill_festa
                font = font_festa

            else:

                if giorno["lavoratori"]:

                    stato = (
                        "👤 "
                        + ", 👤 ".join(
                            giorno[
                                "lavoratori"
                            ]
                        )
                    )

                else:

                    stato = "Tutti in studio"

                fill = (
                    fill_zebra
                    if riga_idx % 2 == 0
                    else PatternFill(
                        fill_type=None
                    )
                )

                font = font_testo

            ws.append([
                nome_sett,
                giorno["data"],
                giorno["giorno"],
                stato
            ])

            ws.row_dimensions[
                riga_idx
            ].height = 22

            for cell in ws[riga_idx]:

                cell.font = font
                cell.alignment = allineamento
                cell.border = bordo

                if fill.fill_type:
                    cell.fill = fill

            riga_idx += 1

    for colonna in ws.columns:

        lunghezza = max(
            (
                len(
                    str(
                        cell.value
                    )
                )
                for cell in colonna
                if cell.value
            ),
            default=12
        )

        ws.column_dimensions[
            get_column_letter(
                colonna[0].column
            )
        ].width = max(
            lunghezza + 4,
            12
        )

    ws.freeze_panes = "A2"

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return send_file(
        output,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=(
            f"calendario_homeworking_"
            f"{start_param}_al_{end_param}.xlsx"
        )
    )


# ============================================================
# TEAM
# ============================================================

@app.route("/team")
def pagina_team():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            nome,
            ruolo
        FROM dipendenti
        ORDER BY id DESC
    """)

    righe = cursor.fetchall()

    cursor.execute("""
        SELECT
            id,
            nome
        FROM ruoli
        ORDER BY nome ASC
    """)

    ruoli = cursor.fetchall()

    cursor.execute("""
        SELECT
            id,
            ruolo,
            max_per_giorno
        FROM limiti_ruoli
        ORDER BY ruolo ASC
    """)

    limiti = {
        r[1]: {
            "id": r[0],
            "max": r[2]
        }
        for r in cursor.fetchall()
    }

    ruoli_unificati = [
        {
            "id": r_id,
            "nome": r_nome,
            "limite_id": limiti.get(
                r_nome,
                {}
            ).get("id"),
            "limite_max": limiti.get(
                r_nome,
                {}
            ).get("max")
        }
        for r_id, r_nome in ruoli
    ]

    conn.close()

    dipendenti = [
        (
            id_dip,
            *split_nome_cognome(
                nome_completo
            ),
            ruolo or "Nessuno",
            nome_completo
        )
        for (
            id_dip,
            nome_completo,
            ruolo
        ) in righe
    ]

    return render_template(
        "team.html",
        dipendenti=dipendenti,
        ruoli=ruoli,
        ruoli_unificati=ruoli_unificati
    )


@app.route("/team/nuovo")
def pagina_nuovo_dipendente():

    conn = get_db()

    ruoli = [
        r[0]
        for r in conn.cursor().execute("""
            SELECT nome
            FROM ruoli
            ORDER BY nome ASC
        """).fetchall()
    ]

    conn.close()

    return render_template(
        "nuovo_dipendente.html",
        ruoli=ruoli
    )


# ============================================================
# DIPENDENTE AGGIUNGI
# ============================================================

@app.route(
    "/dipendente/aggiungi",
    methods=["POST"]
)
def dipendente_aggiungi():

    nome = (
        request.form.get(
            "nome_singolo"
        )
        or request.form.get(
            "nome"
        )
        or ""
    ).strip()

    cognome = (
        request.form.get(
            "cognome_singolo"
        )
        or request.form.get(
            "cognome"
        )
        or ""
    ).strip()

    ruolo = (
        request.form.get(
            "ruolo"
        )
        or "Nessuno"
    ).strip() or "Nessuno"

    if (
        not nome
        or not cognome
        or any(
            char.isdigit()
            for char in (
                nome + cognome
            )
        )
    ):

        return redirect(
            url_for(
                "pagina_team",
                errore="campi_incompleti"
            )
        )

    nome_completo = (
        f"{nome} {cognome}"
    ).strip()

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id
        FROM dipendenti
        WHERE nome = ? COLLATE NOCASE
    """, (
        nome_completo,
    ))

    if cursor.fetchone():

        conn.close()

        return redirect(
            url_for(
                "pagina_team",
                errore="duplicato"
            )
        )

    cursor.execute("""
        SELECT
            COALESCE(
                MAX(rotazione_posizione),
                -1
            ) + 1
        FROM dipendenti
    """)

    posizione = (
        cursor.fetchone()[0]
    )

    cursor.execute("""
        INSERT INTO dipendenti (
            nome,
            giorno_forzato,
            ruolo,
            rotazione_posizione
        )
        VALUES (
            ?,
            'Nessuno',
            ?,
            ?
        )
    """, (
        nome_completo,
        ruolo,
        posizione
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="aggiunto"
        )
    )


# ============================================================
# RUOLO AGGIUNGI
# ============================================================

@app.route(
    "/ruolo/aggiungi",
    methods=["POST"]
)
def ruolo_aggiungi():

    nome = (
        request.form.get(
            "nome_ruolo"
        )
        or ""
    ).strip()

    if not nome:

        return redirect(
            url_for(
                "pagina_team",
                errore="ruolo_vuoto"
            )
        )

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id
        FROM ruoli
        WHERE nome = ? COLLATE NOCASE
    """, (
        nome,
    ))

    if cursor.fetchone():

        conn.close()

        return redirect(
            url_for(
                "pagina_team",
                errore="ruolo_duplicato"
            )
        )

    cursor.execute("""
        INSERT INTO ruoli (nome)
        VALUES (?)
    """, (
        nome,
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="ruolo_aggiunto"
        )
    )


# ============================================================
# RUOLO ELIMINA
# ============================================================

@app.route(
    "/ruolo/elimina/<int:id_ruolo>"
)
def ruolo_elimina(id_ruolo):

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT nome
        FROM ruoli
        WHERE id = ?
    """, (
        id_ruolo,
    ))

    row = cursor.fetchone()

    if row:

        nome_ruolo = row[0]

        cursor.execute("""
            UPDATE dipendenti
            SET ruolo = 'Nessuno'
            WHERE ruolo = ?
        """, (
            nome_ruolo,
        ))

        cursor.execute("""
            DELETE FROM limiti_ruoli
            WHERE ruolo = ?
        """, (
            nome_ruolo,
        ))

        cursor.execute("""
            DELETE FROM ruoli
            WHERE id = ?
        """, (
            id_ruolo,
        ))

        conn.commit()

    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="ruolo_rimosso"
        )
    )


# ============================================================
# DIPENDENTE MODIFICA
# ============================================================

@app.route(
    "/dipendente/aggiorna/<int:id_dipendente>",
    methods=["POST"]
)
def dipendente_aggiorna(
    id_dipendente
):

    nome = (
        request.form.get(
            "nome_singolo"
        )
        or request.form.get(
            "nome"
        )
        or ""
    ).strip()

    cognome = (
        request.form.get(
            "cognome_singolo"
        )
        or request.form.get(
            "cognome"
        )
        or ""
    ).strip()

    ruolo = (
        request.form.get(
            "ruolo"
        )
        or "Nessuno"
    ).strip() or "Nessuno"

    if (
        not nome
        or not cognome
        or any(
            char.isdigit()
            for char in (
                nome + cognome
            )
        )
    ):

        return redirect(
            url_for(
                "pagina_team",
                errore="campi_incompleti"
            )
        )

    nome_completo = (
        f"{nome} {cognome}"
    ).strip()

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id
        FROM dipendenti
        WHERE nome = ? COLLATE NOCASE
          AND id != ?
    """, (
        nome_completo,
        id_dipendente
    ))

    if cursor.fetchone():

        conn.close()

        return redirect(
            url_for(
                "pagina_team",
                errore="duplicato"
            )
        )

    cursor.execute("""
        UPDATE dipendenti
        SET
            nome = ?,
            ruolo = ?
        WHERE id = ?
    """, (
        nome_completo,
        ruolo,
        id_dipendente
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="modificato"
        )
    )


# ============================================================
# DIPENDENTE ELIMINA
# ============================================================

@app.route(
    "/dipendente/elimina/<int:id_dipendente>"
)
def dipendente_elimina(
    id_dipendente
):

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        DELETE FROM eccezioni_dipendenti
        WHERE id_dipendente = ?
    """, (
        id_dipendente,
    ))

    cursor.execute("""
        DELETE FROM dipendenti
        WHERE id = ?
    """, (
        id_dipendente,
    ))

    cursor.execute("""
        SELECT id
        FROM dipendenti
        ORDER BY
            rotazione_posizione ASC,
            id ASC
    """)

    for nuova_posizione, (
        id_corrente,
    ) in enumerate(
        cursor.fetchall()
    ):

        cursor.execute("""
            UPDATE dipendenti
            SET rotazione_posizione = ?
            WHERE id = ?
        """, (
            nuova_posizione,
            id_corrente
        ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="rimosso"
        )
    )


# ============================================================
# TURNI
# ============================================================

@app.route("/turni")
def pagina_turni():

    oggi = datetime.date.today()

    dt_start = datetime.date(
        oggi.year,
        oggi.month,
        1
    )

    dt_end = datetime.date(
        oggi.year,
        oggi.month,
        calendar.monthrange(
            oggi.year,
            oggi.month
        )[1]
    )

    (
        _,
        statistiche,
        tetto_max
    ) = ottieni_calendario_settimanale(
        dt_start,
        dt_end
    )

    conn = get_db()
    cursor = conn.cursor()

    # --------------------------------------------------------
    # FREQUENZA
    # --------------------------------------------------------

    cursor.execute("""
        SELECT valore
        FROM impostazioni
        WHERE chiave = 'frequenza_rotazione'
    """)

    row = cursor.fetchone()

    frequenza_rotazione = (
        normalizza_frequenza(
            row[0]
            if row
            else 1
        )
    )

    # --------------------------------------------------------
    # PROTEZIONE VENERDÌ → LUNEDÌ
    # --------------------------------------------------------

    cursor.execute("""
        SELECT valore
        FROM impostazioni
        WHERE chiave = 'protezione_venerdi_lunedi'
    """)

    row = cursor.fetchone()

    if row is None:

        # Valore predefinito:
        # protezione ATTIVA.
        #
        # In questo modo l'aggiornamento non cambia
        # il comportamento attuale del programma.

        protezione_venerdi_lunedi = True

    else:

        valore = str(
            row[0]
        ).strip().lower()

        protezione_venerdi_lunedi = (
            valore
            in (
                "1",
                "true",
                "yes",
                "on"
            )
        )

    # --------------------------------------------------------
    # TETTO MANUALE
    # --------------------------------------------------------

    cursor.execute("""
        SELECT valore
        FROM impostazioni
        WHERE chiave = 'tetto_manuale'
    """)

    row = cursor.fetchone()

    try:

        tetto_manuale = (
            int(row[0])
            if row
            and row[0]
            not in (None, "")
            else None
        )

    except (
        TypeError,
        ValueError
    ):

        tetto_manuale = None

    if (
        tetto_manuale is not None
        and tetto_manuale < 1
    ):

        tetto_manuale = None

    # --------------------------------------------------------
    # NUMERO AUTOMATICI
    # --------------------------------------------------------

    cursor.execute("""
        SELECT COUNT(*)
        FROM dipendenti
        WHERE
            giorno_forzato = 'Nessuno'
            OR giorno_forzato IS NULL
            OR giorno_forzato = ''
    """)

    n_automatici = (
        cursor.fetchone()[0]
    )

    # --------------------------------------------------------
    # DIPENDENTI
    # --------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            nome,
            giorno_forzato
        FROM dipendenti
        ORDER BY id ASC
    """)

    dipendenti = cursor.fetchall()

    # --------------------------------------------------------
    # ECCEZIONI
    # --------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            id_dipendente,
            data_inizio,
            data_fine
        FROM eccezioni_dipendenti
        ORDER BY
            data_inizio ASC,
            id ASC
    """)

    eccezioni_per_dipendente = {}

    for (
        id_eccezione,
        id_dipendente,
        data_inizio,
        data_fine
    ) in cursor.fetchall():

        di = parse_date(
            data_inizio
        )

        df = parse_date(
            data_fine
        )

        if (
            di
            and df
        ):

            eccezioni_per_dipendente.setdefault(
                id_dipendente,
                []
            ).append(
                (
                    id_eccezione,
                    format_date_it(di),
                    format_date_it(df)
                )
            )

    # --------------------------------------------------------
    # PERIODI
    # --------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            data_inizio,
            data_fine,
            tetto_max
        FROM tetto_periodi
        ORDER BY
            data_inizio ASC,
            id ASC
    """)

    periodi_tetto = []

    for (
        id_periodo,
        data_inizio,
        data_fine,
        tetto
    ) in cursor.fetchall():

        di = parse_date(
            data_inizio
        )

        df = parse_date(
            data_fine
        )

        if (
            di is None
            or df is None
        ):

            continue

        periodi_tetto.append(
            (
                id_periodo,
                format_date_it(di),
                format_date_it(df),
                tetto
            )
        )

    # --------------------------------------------------------
    # LIMITI RUOLO
    # --------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            ruolo,
            max_per_giorno
        FROM limiti_ruoli
        ORDER BY ruolo ASC
    """)

    limiti_ruoli = (
        cursor.fetchall()
    )

    # --------------------------------------------------------
    # RUOLI
    # --------------------------------------------------------

    cursor.execute("""
        SELECT nome
        FROM ruoli
        ORDER BY nome ASC
    """)

    ruoli = [
        r[0]
        for r in cursor.fetchall()
    ]

    conn.close()

    # --------------------------------------------------------
    # TEMPLATE
    # --------------------------------------------------------

    return render_template(
        "turni.html",

        tetto_max=tetto_max,

        tetto_manuale=tetto_manuale,

        frequenza_rotazione=
            frequenza_rotazione,

        # ====================================================
        # NUOVO
        # ====================================================

        protezione_venerdi_lunedi=
            protezione_venerdi_lunedi,

        n_automatici=
            n_automatici,

        dipendenti=
            dipendenti,

        giorni_lista=
            GIORNI_LAVORATIVI,

        statistiche=
            statistiche,

        periodi_tetto=
            periodi_tetto,

        limiti_ruoli=
            limiti_ruoli,

        ruoli=
            ruoli,

        eccezioni_per_dipendente=
            eccezioni_per_dipendente
    )


# ============================================================
# SALVA ROTAZIONE
# ============================================================

@app.route(
    "/turni/salva_rotazione",
    methods=["POST"]
)
def salva_rotazione():

    try:

        settimane = normalizza_frequenza(
            request.form.get(
                "frequenza_rotazione"
            )
        )

        protezione_venerdi_lunedi = (
            request.form.get(
                "protezione_venerdi_lunedi",
                "0"
            ) == "1"
        )

        conn = get_db()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO impostazioni (
                chiave,
                valore
            )
            VALUES (
                'frequenza_rotazione',
                ?
            )
            ON CONFLICT(chiave)
            DO UPDATE SET
                valore = excluded.valore
        """, (
            str(settimane),
        ))

        cursor.execute("""
            INSERT INTO impostazioni (
                chiave,
                valore
            )
            VALUES (
                'protezione_venerdi_lunedi',
                ?
            )
            ON CONFLICT(chiave)
            DO UPDATE SET
                valore = excluded.valore
        """, (
            "1"
            if protezione_venerdi_lunedi
            else "0",
        ))

        conn.commit()
        conn.close()

        return redirect(
            url_for(
                "pagina_turni",
                salvato="rotazione"
            )
        )

    except Exception as e:

        print(
            "Errore salva_rotazione:",
            e
        )

        return redirect(
            url_for(
                "pagina_turni",
                errore="rotazione"
            )
        )


# ============================================================
# SALVA TUTTE LE IMPOSTAZIONI
# ============================================================

@app.route(
    "/turni/salva_impostazioni",
    methods=["POST"]
)
def salva_impostazioni():

    conn = None

    try:

        # ----------------------------------------------------
        # SETTIMANE
        # ----------------------------------------------------

        settimane = normalizza_frequenza(
            request.form.get(
                "frequenza_rotazione",
                "1"
            )
        )

        # ----------------------------------------------------
        # PERSONE HOMEWORKING AL GIORNO
        # ----------------------------------------------------

        try:

            persone_giorno = int(
                request.form.get(
                    "tetto_manuale",
                    "1"
                )
            )

        except (
            TypeError,
            ValueError
        ):

            raise ValueError(
                "Numero persone non valido"
            )

        if persone_giorno < 1:
            raise ValueError(
                "Il tetto deve essere almeno 1"
            )

        # ----------------------------------------------------
        # VENERDÌ → LUNEDÌ
        # ----------------------------------------------------

        protezione_venerdi_lunedi = (
            request.form.get(
                "protezione_venerdi_lunedi",
                "0"
            ) == "1"
        )

        # ----------------------------------------------------
        # DATABASE
        # ----------------------------------------------------

        conn = get_db()
        cursor = conn.cursor()

        # ----------------------------------------------------
        # SALVA SETTIMANE
        # ----------------------------------------------------

        cursor.execute("""
            INSERT INTO impostazioni (
                chiave,
                valore
            )
            VALUES (
                'frequenza_rotazione',
                ?
            )
            ON CONFLICT(chiave)
            DO UPDATE SET
                valore = excluded.valore
        """, (
            str(settimane),
        ))

        # ----------------------------------------------------
        # SALVA PERSONE/GIORNO
        # ----------------------------------------------------

        cursor.execute("""
            INSERT INTO impostazioni (
                chiave,
                valore
            )
            VALUES (
                'tetto_manuale',
                ?
            )
            ON CONFLICT(chiave)
            DO UPDATE SET
                valore = excluded.valore
        """, (
            str(persone_giorno),
        ))

        # ----------------------------------------------------
        # SALVA VENERDÌ → LUNEDÌ
        # ----------------------------------------------------

        cursor.execute("""
            INSERT INTO impostazioni (
                chiave,
                valore
            )
            VALUES (
                'protezione_venerdi_lunedi',
                ?
            )
            ON CONFLICT(chiave)
            DO UPDATE SET
                valore = excluded.valore
        """, (
            "1"
            if protezione_venerdi_lunedi
            else "0",
        ))

        # ----------------------------------------------------
        # UNICO COMMIT
        # ----------------------------------------------------

        conn.commit()
        conn.close()
        conn = None

        return redirect(
            url_for(
                "pagina_turni",
                salvato="impostazioni"
            )
        )

    except Exception as e:

        print(
            "Errore salva_impostazioni:",
            e
        )

        if conn is not None:

            try:
                conn.rollback()
                conn.close()

            except Exception:
                pass

        return redirect(
            url_for(
                "pagina_turni",
                errore="impostazioni"
            )
        )


# ============================================================
# SALVA TETTO
# ============================================================

@app.route(
    "/turni/salva_tetto",
    methods=["POST"]
)
def salva_tetto():

    azione = request.form.get(
        "azione",
        "salva"
    )

    conn = get_db()
    cursor = conn.cursor()

    # --------------------------------------------------------
    # MODALITÀ AUTOMATICA
    # --------------------------------------------------------

    if azione == "auto":

        cursor.execute("""
            DELETE FROM impostazioni
            WHERE chiave = 'tetto_manuale'
        """)

        conn.commit()
        conn.close()

        return redirect(
            url_for(
                "pagina_turni",
                salvato="tetto"
            )
        )

    # --------------------------------------------------------
    # LETTURA TETTO
    # --------------------------------------------------------

    try:

        persone_giorno = int(
            request.form.get(
                "tetto_manuale"
            )
        )

        if persone_giorno < 1:
            raise ValueError

    except (
        TypeError,
        ValueError
    ):

        conn.close()

        return redirect(
            url_for(
                "pagina_turni",
                errore="tetto"
            )
        )

    # --------------------------------------------------------
    # SALVA SOLO IL TETTO
    #
    # NON MODIFICA PIÙ LE SETTIMANE
    # --------------------------------------------------------

    cursor.execute("""
        INSERT INTO impostazioni (
            chiave,
            valore
        )
        VALUES (
            'tetto_manuale',
            ?
        )
        ON CONFLICT(chiave)
        DO UPDATE SET
            valore = excluded.valore
    """, (
        str(
            persone_giorno
        ),
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_turni",
            salvato="tetto"
        )
    )


# ============================================================
# FORZA PER PERIODO
# ============================================================

@app.route(
    "/turni/periodo/aggiungi",
    methods=["POST"]
)
def periodo_aggiungi():

    data_inizio = parse_date(
        request.form.get(
            "data_inizio"
        )
    )

    data_fine = parse_date(
        request.form.get(
            "data_fine"
        )
    )

    try:

        tetto_periodo = int(
            request.form.get(
                "tetto_periodo"
            )
        )

    except (
        TypeError,
        ValueError
    ):

        tetto_periodo = 0

    # --------------------------------------------------------
    # VALIDAZIONE
    # --------------------------------------------------------

    if data_inizio is None:
        return redirect(
            url_for(
                "pagina_turni",
                errore="periodo"
            )
        )

    if data_fine is None:
        return redirect(
            url_for(
                "pagina_turni",
                errore="periodo"
            )
        )

    if data_fine < data_inizio:
        return redirect(
            url_for(
                "pagina_turni",
                errore="periodo"
            )
        )

    if tetto_periodo < 1:
        return redirect(
            url_for(
                "pagina_turni",
                errore="periodo"
            )
        )

    conn = get_db()
    cursor = conn.cursor()

    # --------------------------------------------------------
    # SALVATAGGIO
    # --------------------------------------------------------

    cursor.execute("""
        INSERT INTO tetto_periodi (
            data_inizio,
            data_fine,
            tetto_max
        )
        VALUES (?, ?, ?)
    """, (
        data_inizio.strftime(
            "%Y-%m-%d"
        ),
        data_fine.strftime(
            "%Y-%m-%d"
        ),
        tetto_periodo
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_turni",
            salvato="periodo"
        )
    )


# ============================================================
# ELIMINA PERIODO
# ============================================================

@app.route(
    "/turni/periodo/elimina/<int:id_periodo>"
)
def periodo_elimina(
    id_periodo
):

    conn = get_db()

    conn.execute("""
        DELETE FROM tetto_periodi
        WHERE id = ?
    """, (
        id_periodo,
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_turni"
        )
    )


# ============================================================
# LIMITE RUOLO
# ============================================================

@app.route(
    "/turni/ruolo_limite/aggiungi",
    methods=["POST"]
)
def ruolo_limite_aggiungi():

    ruolo = (
        request.form.get(
            "ruolo"
        )
        or ""
    ).strip()

    try:

        massimo = int(
            request.form.get(
                "max_per_giorno"
            )
        )

    except (
        TypeError,
        ValueError
    ):

        massimo = None

    if (
        not ruolo
        or massimo is None
        or massimo < 0
    ):

        return redirect(
            url_for(
                "pagina_team",
                errore="limite_ruolo"
            )
        )

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        INSERT INTO limiti_ruoli (
            ruolo,
            max_per_giorno
        )
        VALUES (?, ?)
        ON CONFLICT(ruolo)
        DO UPDATE SET
            max_per_giorno = excluded.max_per_giorno
    """, (
        ruolo,
        massimo
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="limite_ruolo"
        )
    )


# ============================================================
# ELIMINA LIMITE RUOLO
# ============================================================

@app.route(
    "/turni/ruolo_limite/elimina/<int:id_limite>"
)
def ruolo_limite_elimina(
    id_limite
):

    conn = get_db()

    conn.execute("""
        DELETE FROM limiti_ruoli
        WHERE id = ?
    """, (
        id_limite,
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_team",
            salvato="limite_ruolo_rimosso"
        )
    )


# ============================================================
# ECCEZIONE AGGIUNGI
# ============================================================

@app.route(
    "/turni/eccezione/aggiungi",
    methods=["POST"]
)
def eccezione_aggiungi():

    try:

        id_dipendente = int(
            request.form.get(
                "id_dipendente"
            )
        )

    except (
        TypeError,
        ValueError
    ):

        return redirect(
            url_for(
                "pagina_turni",
                errore="eccezione"
            )
        )

    data_inizio = parse_date(
        request.form.get(
            "eccezione_data_inizio"
        )
    )

    data_fine = parse_date(
        request.form.get(
            "eccezione_data_fine"
        )
    )

    if (
        data_inizio is None
        or data_fine is None
        or data_fine < data_inizio
    ):

        return redirect(
            url_for(
                "pagina_turni",
                errore="eccezione"
            )
        )

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id
        FROM dipendenti
        WHERE id = ?
    """, (
        id_dipendente,
    ))

    if cursor.fetchone() is None:

        conn.close()

        return redirect(
            url_for(
                "pagina_turni",
                errore="eccezione"
            )
        )

    cursor.execute("""
        INSERT INTO eccezioni_dipendenti (
            id_dipendente,
            data_inizio,
            data_fine
        )
        VALUES (?, ?, ?)
    """, (
        id_dipendente,
        data_inizio.strftime(
            "%Y-%m-%d"
        ),
        data_fine.strftime(
            "%Y-%m-%d"
        )
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_turni",
            salvato="eccezione"
        )
    )


# ============================================================
# ECCEZIONE ELIMINA
# ============================================================

@app.route(
    "/turni/eccezione/elimina/<int:id_eccezione>"
)
def eccezione_elimina(
    id_eccezione
):

    conn = get_db()

    conn.execute("""
        DELETE FROM eccezioni_dipendenti
        WHERE id = ?
    """, (
        id_eccezione,
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_turni",
            salvato="eccezione"
        )
    )


# ============================================================
# SALVA TURNI INDIVIDUALI
# ============================================================

@app.route(
    "/turni/salva_singoli",
    methods=["POST"]
)
def salva_singoli():

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id
        FROM dipendenti
        ORDER BY id ASC
    """)

    for (
        id_dipendente,
    ) in cursor.fetchall():

        giorno = request.form.get(
            f"giorno_forzato_{id_dipendente}"
        )

        if giorno in (
            None,
            ""
        ):
            continue

        if (
            giorno != "Nessuno"
            and giorno
            not in GIORNI_LAVORATIVI
        ):
            continue

        cursor.execute("""
            UPDATE dipendenti
            SET giorno_forzato = ?
            WHERE id = ?
        """, (
            giorno,
            id_dipendente
        ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_turni",
            salvato="turni"
        )
    )


# ============================================================
# FESTIVITÀ
# ============================================================

@app.route("/festivita")
def pagina_festivita():

    conn = get_db()

    righe = conn.cursor().execute("""
        SELECT
            id,
            data,
            descrizione
        FROM festivita
        ORDER BY data DESC
    """).fetchall()

    conn.close()

    festivita = []

    for (
        id_festa,
        data_raw,
        descrizione
    ) in righe:

        dt = parse_date(
            data_raw
        )

        if dt:

            festivita.append(
                (
                    id_festa,
                    format_date_it(dt),
                    descrizione
                )
            )

    return render_template(
        "feste.html",
        festivita=festivita
    )


# ============================================================
# FESTIVITÀ AGGIUNGI
# ============================================================

@app.route(
    "/festa/aggiungi",
    methods=["POST"]
)
def festa_aggiungi():

    data = parse_date(
        request.form.get(
            "data"
        )
    )

    descrizione = (
        request.form.get(
            "descrizione"
        )
        or ""
    ).strip()

    if (
        data is None
        or not descrizione
    ):

        return redirect(
            url_for(
                "pagina_festivita",
                errore="data"
            )
        )

    conn = get_db()
    cursor = conn.cursor()

    cursor.execute("""
        INSERT INTO festivita (
            data,
            descrizione
        )
        VALUES (?, ?)
        ON CONFLICT(data)
        DO UPDATE SET
            descrizione = excluded.descrizione
    """, (
        data.strftime(
            "%Y-%m-%d"
        ),
        descrizione
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_festivita",
            salvato="aggiunto"
        )
    )


# ============================================================
# FESTIVITÀ ELIMINA
# ============================================================

@app.route(
    "/festa/elimina/<int:id_festa>"
)
def festa_elimina(id_festa):

    conn = get_db()

    conn.execute("""
        DELETE FROM festivita
        WHERE id = ?
    """, (
        id_festa
    ))

    conn.commit()
    conn.close()

    return redirect(
        url_for(
            "pagina_festivita",
            salvato="rimosso"
        )
    )


# ============================================================
# AVVIO
# ============================================================

if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )
